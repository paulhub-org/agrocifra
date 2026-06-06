"""Демонстрационные пользователи, справочники и пилотные данные (Спринты 4, 8).

Запуск: python -m app.db.seed   (после `alembic upgrade head`).
Пароли демонстрационные — сменить перед реальной эксплуатацией.
"""
from pathlib import Path

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.db.session import SessionLocal
from app.models.models import Organization, Region, Role, UserAccount
from app.services import etl
from app.services.normalize import clean_text

DEMO_USERS = [
    ("org", "org123", Role.organization, "Оператор организации"),
    ("region", "region123", Role.regional_operator, "Региональный оператор"),
    ("office", "office123", Role.digitalization_office, "Офис цифровизации"),
    ("gov", "gov123", Role.state_authority, "Государственный орган"),
]
DEMO_REGIONS = ["Брестская", "Витебская", "Гомельская", "Гродненская", "Минская", "Могилёвская"]

SEED_DATA = Path(__file__).parent / "seed_data"
PILOT_FILE = SEED_DATA / "Pilot_Organizations_v2_0.xlsx"
DEBT_FILE = SEED_DATA / "оценка_проекта_и_долгового_риска.xlsx"
DEBT_ORG = "РСДУП «Шипяны-АСК»"  # модель долгового риска — единичный проект; привязываем к ведущей организации

# Сопоставление текста адреса с нормализованным наименованием области
_OBLAST = {
    "Брестская": "Брестская", "Витебская": "Витебская", "Гомельская": "Гомельская",
    "Гродненская": "Гродненская", "Минская": "Минская",
    "Могилевская": "Могилёвская", "Могилёвская": "Могилёвская",
}


def _region_name(address: str | None) -> str | None:
    for key, name in _OBLAST.items():
        if key in (address or ""):
            return name
    return None


def _get_or_create_region(db: Session, name: str) -> Region:
    reg = db.execute(select(Region).where(Region.name == name)).scalar_one_or_none()
    if not reg:
        reg = Region(name=name)
        db.add(reg)
        db.flush()
    return reg


def _assign_regions_and_addresses(db: Session) -> None:
    """Прочитать строку адреса пилотного файла и проставить организациям область и адрес."""
    wb = openpyxl.load_workbook(PILOT_FILE, data_only=True)
    ws = wb["Цифровая зрелость_data"]
    for col in range(2, ws.max_column + 1):
        name = clean_text(ws.cell(1, col).value)
        address = clean_text(ws.cell(2, col).value)
        if not name:
            continue
        org = db.execute(select(Organization).where(Organization.name == name)).scalar_one_or_none()
        if not org:
            continue
        if address and not org.address:
            org.address = address
        region_name = _region_name(address)
        if region_name:
            org.region_id = _get_or_create_region(db, region_name).id
    db.commit()


def seed_pilot(db: Session) -> None:
    """Загрузить пилотные организации (зрелость, эффективность) и модель долгового риска."""
    if PILOT_FILE.exists():
        etl.import_maturity_from_excel(db, PILOT_FILE)
        etl.import_efficiency_from_excel(db, PILOT_FILE)
        _assign_regions_and_addresses(db)
    if DEBT_FILE.exists():
        try:
            etl.load_project_finance(db, DEBT_ORG, DEBT_FILE)
        except Exception:  # noqa: BLE001 — сидирование финмодели не должно прерывать запуск
            pass


def link_demo_users(db: Session) -> None:
    """Привязать демо-роли к данным: «org» → ведущая организация, «region» → Минская область."""
    org = db.execute(select(Organization).where(Organization.name.like("%Шипяны%"))).scalar_one_or_none()
    if org:
        u = db.execute(select(UserAccount).where(UserAccount.login == "org")).scalar_one_or_none()
        if u and not u.organization_id:
            u.organization_id = org.id
            if org.region_id:
                u.region_id = org.region_id
    minsk = db.execute(select(Region).where(Region.name == "Минская")).scalar_one_or_none()
    if minsk:
        u = db.execute(select(UserAccount).where(UserAccount.login == "region")).scalar_one_or_none()
        if u and not u.region_id:
            u.region_id = minsk.id
    db.commit()


def seed(db: Session) -> None:
    for name in DEMO_REGIONS:
        _get_or_create_region(db, name)
    for login, pw, role, full_name in DEMO_USERS:
        if not db.execute(select(UserAccount).where(UserAccount.login == login)).scalar_one_or_none():
            db.add(UserAccount(login=login, password_hash=hash_password(pw),
                               role=role.value, full_name=full_name, is_active=True))
    db.commit()


if __name__ == "__main__":
    db = SessionLocal()
    try:
        seed(db)            # пользователи и справочник областей (детерминированно)
        seed_pilot(db)      # пилотные организации, оценки и модель долгового риска
        link_demo_users(db) # привязка демо-ролей к организации/области
        print("Демо-данные и пилотные организации созданы. Пользователи: org / region / office / gov (пароли *123).")
    finally:
        db.close()
