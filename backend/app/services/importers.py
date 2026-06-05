"""Импорт исходных данных из Excel/CSV.

Поддерживает структуру книги Pilot_Organizations_v2_0.xlsx (листы
«Цифровая зрелость_data» и «Эффективность_data», по 7 организаций), а также
произвольный CSV. «Урожайность ПОСЛЕ» вычисляется как валовый сбор ÷ площадь
обрабатываемых земель (отдельным полем в форме не хранится).
"""
from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

from app.services.efficiency import EfficiencyIndexInputs
from app.services.normalize import clean_text, parse_decimal

SHEET_MATURITY = "Цифровая зрелость_data"
SHEET_EFFICIENCY = "Эффективность_data"

# Лист зрелости: номера строк (1-based) с агрегатами и итоговым уровнем.
MATURITY_ROWS = {"name": 1, "need_avg": 9, "capability_avg": 14, "reference_maturity": 15}

# Лист эффективности: поле входа → номер строки (раскладка книги v2).
EFFICIENCY_ROW_MAP: dict[str, int] = {
    "cost_total_before": 28, "area_before_ha": 30, "cost_per_ha_after": 26,
    "yield_before": 16,
    "profit_after": 47, "total_costs": 4, "profit_before": 46, "total_costs_before": 5,
    "petrol_before_t": 34, "petrol_after_t": 35, "diesel_before_t": 36, "diesel_after_t": 37,
    "productivity_after": 10, "productivity_before": 3, "taxes_after": 53, "taxes_before": 52,
}
# «Урожайность ПОСЛЕ» = валовый сбор (r25) / площадь обрабатываемых земель (r24)
GROSS_HARVEST_ROW = 25
AREA_CURRENT_ROW = 24


def _org_columns(ws, name_row: int) -> list[int]:
    cols: list[int] = []
    c = 2
    while ws.cell(name_row, c).value not in (None, ""):
        cols.append(c)
        c += 1
    return cols


def import_maturity_sheet(path: str | Path, sheet: str = SHEET_MATURITY) -> list[dict]:
    """Прочитать лист зрелости → [{name, need_avg, capability_avg, reference_maturity}]."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    records: list[dict] = []
    for col in _org_columns(ws, MATURITY_ROWS["name"]):
        name = clean_text(ws.cell(MATURITY_ROWS["name"], col).value)
        if not name:
            continue
        records.append({
            "name": name,
            "need_avg": parse_decimal(ws.cell(MATURITY_ROWS["need_avg"], col).value),
            "capability_avg": parse_decimal(ws.cell(MATURITY_ROWS["capability_avg"], col).value),
            "reference_maturity": parse_decimal(ws.cell(MATURITY_ROWS["reference_maturity"], col).value),
        })
    return records


def import_efficiency_sheet(path: str | Path, sheet: str = SHEET_EFFICIENCY,
                            row_map: dict[str, int] | None = None) -> list[dict]:
    """Прочитать лист эффективности → [{name, <нормализованные поля>}].

    «Урожайность ПОСЛЕ» вычисляется из валового сбора и площади обрабатываемых земель.
    """
    row_map = row_map or EFFICIENCY_ROW_MAP
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    records: list[dict] = []
    for col in _org_columns(ws, 1):
        name = clean_text(ws.cell(1, col).value)
        if not name:
            continue
        rec: dict = {"name": name}
        for field, row in row_map.items():
            rec[field] = parse_decimal(ws.cell(row, col).value)
        gross = parse_decimal(ws.cell(GROSS_HARVEST_ROW, col).value)
        area = parse_decimal(ws.cell(AREA_CURRENT_ROW, col).value)
        if gross is not None and area:
            rec["yield_after"] = gross / area
        records.append(rec)
    return records


def to_efficiency_inputs(rec: dict) -> EfficiencyIndexInputs:
    """Построить EfficiencyIndexInputs из нормализованной записи (требует yield_after)."""
    required = ["cost_total_before", "area_before_ha", "cost_per_ha_after",
                "yield_before", "yield_after", "petrol_before_t", "petrol_after_t",
                "diesel_before_t", "diesel_after_t", "productivity_after",
                "productivity_before", "taxes_after", "taxes_before"]
    missing = [k for k in required if rec.get(k) is None]
    if missing:
        raise ValueError(f"{rec.get('name','?')}: отсутствуют поля {', '.join(missing)}")
    fields = {k: rec.get(k) for k in EfficiencyIndexInputs.__dataclass_fields__}
    return EfficiencyIndexInputs(**fields)


def import_csv(path: str | Path) -> list[dict]:
    """Прочитать CSV (заголовок = имена полей) → [{поле: нормализованное значение}]."""
    records: list[dict] = []
    with open(path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            rec: dict = {}
            for k, v in row.items():
                if k is None:
                    continue
                key = k.strip()
                rec[key] = clean_text(v) if key in ("name", "region") else parse_decimal(v)
            records.append(rec)
    return records
