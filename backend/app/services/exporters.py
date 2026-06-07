"""Модуль экспорта отчётов в Word, Excel и PDF (Спринт 5).

Форматирование ориентировано на приложения к диссертации. Каждая функция
возвращает (bytes, имя_файла, media_type) для отдачи через StreamingResponse.
"""
from __future__ import annotations

import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.services.reporting import summary

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"

_GREEN = "2D6A4F"
_HEADER_FILL = PatternFill("solid", fgColor="D8F3DC")


def _fmt(x) -> str:
    return "—" if x is None else f"{x:.2f}".replace(".", ",")


# ─────────────────────────────── EXCEL ───────────────────────────────
def build_excel(db: Session) -> tuple[bytes, str, str]:
    data = summary(db)
    wb = Workbook()

    ws = wb.active
    ws.title = "Эффективность"
    eff_head = ["Организация", "Регион", "Экономический", "Экологический",
                "Социальный", "КЭц", "Зона"]
    ws.append(eff_head)
    for row in data["efficiency_by_org"]:
        ws.append([row["name"], row["region"], row["economic_index"], row["ecological_index"],
                   row["social_index"], row["coefficient"], row["zone"]])

    ws2 = wb.create_sheet("Зрелость")
    ws2.append(["Организация", "Регион", "Потребность", "Возможности",
                "Цифровая зрелость", "Зона"])
    for row in data["maturity_by_org"]:
        ws2.append([row["name"], row["region"], row["need_avg"], row["capability_avg"],
                    row["maturity"], row["zone"]])

    ws3 = wb.create_sheet("Сводка по регионам")
    ws3.append(["Регион", "Среднее КЭц", "Оценок эффективности",
                "Средняя зрелость", "Оценок зрелости"])
    for row in data["by_region"]:
        ws3.append([row["region"], row["mean_ke"], row["efficiency_count"],
                    row["mean_maturity"], row["maturity_count"]])

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color=_GREEN)
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col_cells in sheet.columns:
            width = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
            sheet.column_dimensions[col_cells[0].column_letter].width = min(max(width + 2, 12), 42)
        sheet.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "АгроЦифра_отчёт.xlsx", XLSX_MIME


# ─────────────────────────────── WORD ───────────────────────────────
def build_word(db: Session) -> tuple[bytes, str, str]:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    from docx.shared import Cm

    data = summary(db)
    doc = Document()
    # Поля страницы (Инструкция ВАК / требования заказчика): верх/низ 2 см, левое 3 см, правое 1 см
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1)
    # Автор документа — отсутствует (по умолчанию python-docx проставляет «python-docx»)
    cp = doc.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY  # выравнивание по ширине

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    h.add_run("ПРИЛОЖЕНИЕ").bold = True
    title = doc.add_heading(
        "Результаты оценки эффективности и цифровой зрелости организаций", level=1
    )
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1B, 0x43, 0x32)

    o = data["overall"]
    doc.add_paragraph(
        f"Всего организаций: {o['organizations']}. Оценок эффективности: "
        f"{o['efficiency_count']} (эффективных, КЭц > 1,0: {o['efficiency_effective']}), "
        f"среднее значение КЭц: {_fmt(o['mean_ke'])}. Оценок цифровой зрелости: "
        f"{o['maturity_count']}, средний уровень: {_fmt(o['mean_maturity'])}."
    )

    doc.add_heading("Таблица 1 — Эффективность цифровизации по организациям", level=2)
    t = doc.add_table(rows=1, cols=5)
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(["Организация", "Регион", "КЭц", "Зона", "Эк./Экол./Соц."]):
        t.rows[0].cells[i].paragraphs[0].add_run(c).bold = True
    for row in data["efficiency_by_org"]:
        cells = t.add_row().cells
        cells[0].text = row["name"]
        cells[1].text = row["region"]
        cells[2].text = _fmt(row["coefficient"])
        cells[3].text = row["zone"]
        cells[4].text = f"{_fmt(row['economic_index'])} / {_fmt(row['ecological_index'])} / {_fmt(row['social_index'])}"

    doc.add_heading("Таблица 2 — Цифровая зрелость по организациям", level=2)
    t2 = doc.add_table(rows=1, cols=5)
    t2.style = "Light Grid Accent 1"
    for i, c in enumerate(["Организация", "Регион", "Потребность", "Возможности", "Зрелость (зона)"]):
        t2.rows[0].cells[i].paragraphs[0].add_run(c).bold = True
    for row in data["maturity_by_org"]:
        cells = t2.add_row().cells
        cells[0].text = row["name"]
        cells[1].text = row["region"]
        cells[2].text = _fmt(row["need_avg"])
        cells[3].text = _fmt(row["capability_avg"])
        cells[4].text = f"{_fmt(row['maturity'])} ({row['zone']})"

    if data["by_region"]:
        doc.add_heading("Таблица 3 — Сводка по регионам", level=2)
        t3 = doc.add_table(rows=1, cols=3)
        t3.style = "Light Grid Accent 1"
        for i, c in enumerate(["Регион", "Среднее КЭц", "Средняя зрелость"]):
            t3.rows[0].cells[i].paragraphs[0].add_run(c).bold = True
        for row in data["by_region"]:
            cells = t3.add_row().cells
            cells[0].text = row["region"]
            cells[1].text = _fmt(row["mean_ke"])
            cells[2].text = _fmt(row["mean_maturity"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), "АгроЦифра_отчёт.docx", DOCX_MIME


# ─────────────────────────────── PDF ───────────────────────────────
def _register_cyrillic_font() -> tuple[str, str]:
    """Зарегистрировать кириллический шрифт для PDF.

    Шрифты DejaVu поставляются в составе приложения (app/services/fonts), что
    гарантирует корректный экспорт PDF в любом окружении, включая slim-образ
    Docker, где системные шрифты отсутствуют. Системный путь — запасной вариант.
    """
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    bundled = Path(__file__).resolve().parent / "fonts"
    system = Path("/usr/share/fonts/truetype/dejavu")
    for base in (bundled, system):
        regular = base / "DejaVuSerif.ttf"
        bold = base / "DejaVuSerif-Bold.ttf"
        if regular.exists() and bold.exists():
            if "DejaVuSerif" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("DejaVuSerif", str(regular)))
                pdfmetrics.registerFont(TTFont("DejaVuSerif-Bold", str(bold)))
            return "DejaVuSerif", "DejaVuSerif-Bold"
    raise RuntimeError("Не найден кириллический шрифт DejaVuSerif для экспорта в PDF")


def build_pdf(db: Session) -> tuple[bytes, str, str]:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    font, font_bold = _register_cyrillic_font()
    data = summary(db)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm,
                            leftMargin=2 * cm, rightMargin=2 * cm)
    styles = getSampleStyleSheet()
    styles["Normal"].fontName = font
    styles["Title"].fontName = font_bold
    styles["Title"].textColor = colors.HexColor("#1B4332")
    styles["Heading2"].fontName = font_bold
    styles["Heading2"].textColor = colors.HexColor("#2D6A4F")

    def make_table(headers, rows):
        table = Table([headers] + rows, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font),
            ("FONTNAME", (0, 0), (-1, 0), font_bold),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D8F3DC")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1B4332")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D6CC")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6FAF7")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return table

    o = data["overall"]
    elems = [
        Paragraph("Результаты оценки эффективности и цифровой зрелости организаций", styles["Title"]),
        Spacer(1, 6),
        Paragraph(
            f"Всего организаций: {o['organizations']}. Среднее КЭц: {_fmt(o['mean_ke'])} "
            f"(эффективных: {o['efficiency_effective']} из {o['efficiency_count']}). "
            f"Средний уровень цифровой зрелости: {_fmt(o['mean_maturity'])}.",
            styles["Normal"]),
        Spacer(1, 12),
        Paragraph("Таблица 1 — Эффективность цифровизации по организациям", styles["Heading2"]),
    ]
    elems.append(make_table(
        ["Организация", "Регион", "КЭц", "Зона"],
        [[r["name"], r["region"], _fmt(r["coefficient"]), r["zone"]]
         for r in data["efficiency_by_org"]] or [["—", "—", "—", "—"]]))
    elems += [Spacer(1, 14),
              Paragraph("Таблица 2 — Цифровая зрелость по организациям", styles["Heading2"])]
    elems.append(make_table(
        ["Организация", "Регион", "Зрелость", "Зона"],
        [[r["name"], r["region"], _fmt(r["maturity"]), r["zone"]]
         for r in data["maturity_by_org"]] or [["—", "—", "—", "—"]]))
    if data["by_region"]:
        elems += [Spacer(1, 14), Paragraph("Таблица 3 — Сводка по регионам", styles["Heading2"])]
        elems.append(make_table(
            ["Регион", "Среднее КЭц", "Средняя зрелость"],
            [[r["region"], _fmt(r["mean_ke"]), _fmt(r["mean_maturity"])] for r in data["by_region"]]))

    doc.build(elems)
    return buf.getvalue(), "АгроЦифра_отчёт.pdf", PDF_MIME


# ──────────────────────── РЕКОМЕНДАЦИИ (V2.0, задача 3) ────────────────────────
def _rec_filename(ext: str) -> str:
    return f"АгроЦифра_рекомендации.{ext}"


def _group_recs(items: list[dict]) -> list[dict]:
    """Сгруппировать рекомендации по организациям, сохраняя порядок поступления."""
    groups: list[dict] = []
    index: dict = {}
    for it in items:
        oid = it["organization_id"]
        if oid not in index:
            index[oid] = {"organization": it["organization"], "region": it["region"],
                          "district": it["district"], "items": []}
            groups.append(index[oid])
        index[oid]["items"].append(it)
    return groups


def _loc(group: dict) -> str:
    bits = [b for b in (group.get("region"), group.get("district")) if b and b != "Не указан"]
    return ", ".join(bits)


def _rec_header(it: dict) -> str:
    """Заголовок рекомендации: «Показатель — уровень (значение).»"""
    head = it["indicator_label"]
    if it.get("level"):
        head += f" — {it['level']}"
    if it.get("value") is not None:
        head += f" ({_fmt(it['value'])})"
    return head + "."


def build_recommendations_excel(items: list[dict]) -> tuple[bytes, str, str]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Рекомендации"
    head = ["Организация", "Область", "Район", "Показатель", "Значение", "Уровень/зона", "Рекомендация"]
    ws.append(head)
    for it in items:
        ws.append([it["organization"], it.get("region") or "", it.get("district") or "",
                   it["indicator_label"], _fmt(it.get("value")), it.get("level") or "—", it["text"]])
    widths = [30, 14, 16, 32, 10, 14, 90]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True, color=_GREEN)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        row[6].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), _rec_filename("xlsx"), XLSX_MIME


def build_recommendations_word(items: list[dict]) -> tuple[bytes, str, str]:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1)
    cp = doc.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY  # выравнивание по ширине

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    h.add_run("ПРИЛОЖЕНИЕ").bold = True
    title = doc.add_heading("Рекомендации по повышению показателей цифровизации", level=1)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1B, 0x43, 0x32)

    groups = _group_recs(items)
    if not groups:
        doc.add_paragraph("По заданным условиям рекомендации отсутствуют.")
    for g in groups:
        loc = _loc(g)
        doc.add_heading(g["organization"] + (f" ({loc})" if loc else ""), level=2)
        for it in g["items"]:
            p = doc.add_paragraph()
            p.add_run(_rec_header(it) + " ").bold = True
            p.add_run(it["text"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), _rec_filename("docx"), DOCX_MIME


def build_recommendations_pdf(items: list[dict]) -> tuple[bytes, str, str]:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    font, font_bold = _register_cyrillic_font()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm,
                            leftMargin=2 * cm, rightMargin=2 * cm)
    styles = getSampleStyleSheet()
    styles["Normal"].fontName = font
    styles["Normal"].fontSize = 10
    styles["Normal"].leading = 14
    styles["Title"].fontName = font_bold
    styles["Title"].textColor = colors.HexColor("#1B4332")
    styles["Heading2"].fontName = font_bold
    styles["Heading2"].textColor = colors.HexColor("#2D6A4F")

    def esc(s: str) -> str:
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    elems = [Paragraph("Рекомендации по повышению показателей цифровизации", styles["Title"]),
             Spacer(1, 10)]
    groups = _group_recs(items)
    if not groups:
        elems.append(Paragraph("По заданным условиям рекомендации отсутствуют.", styles["Normal"]))
    for g in groups:
        loc = _loc(g)
        elems.append(Paragraph(esc(g["organization"] + (f" ({loc})" if loc else "")), styles["Heading2"]))
        for it in g["items"]:
            elems.append(Paragraph(f"<b>{esc(_rec_header(it))}</b> {esc(it['text'])}", styles["Normal"]))
            elems.append(Spacer(1, 4))
        elems.append(Spacer(1, 8))
    doc.build(elems)
    return buf.getvalue(), _rec_filename("pdf"), PDF_MIME
