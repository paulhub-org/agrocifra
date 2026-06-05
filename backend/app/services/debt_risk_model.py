"""Подключение Excel-модели «Оценка проекта и долгового риска».

Разбирает заполненный шаблон модели и извлекает стоимость (CAPEX), эффект (NPV),
доходность (IRR) и показатели долгового риска (DSCR/Кпз, ICR), а также стресс-оценки.
На основе этих данных формируется кандидат `OptProject` для модуля оптимизации:
стоимость = суммарный CAPEX, эффект = NPV, предел финансирования (credit_limit) —
по запасу обслуживания долга относительно норматива DSCR (граница кредитоспособности).

Числовые значения берутся из кэшированных результатов формул листа
«3. Исх проект данные и риски»; строки находятся по тексту показателя (столбец B),
что устойчиво к смещению строк шаблона.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from app.services.optimization import OptProject

# Норматив коэффициента покрытия задолженности (Кпз/DSCR) — постановление
# Минэкономики РБ № 158; используется как граница кредитоспособности.
DSCR_NORM_DEFAULT = 1.3


@dataclass
class ProjectFinance:
    capex_by_year: list[float] = field(default_factory=list)
    capex_total: float = 0.0
    npv: float | None = None
    irr: float | None = None
    npv_stress: float | None = None
    irr_stress: float | None = None
    dscr_by_year: list[float] = field(default_factory=list)
    dscr_min: float | None = None
    dscr_stress_min: float | None = None
    icr_by_year: list[float] = field(default_factory=list)
    icr_min: float | None = None
    cash_flows: list[float] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "capex_total": self.capex_total, "capex_by_year": self.capex_by_year,
            "npv": self.npv, "irr": self.irr,
            "npv_stress": self.npv_stress, "irr_stress": self.irr_stress,
            "dscr_min": self.dscr_min, "dscr_stress_min": self.dscr_stress_min,
            "icr_min": self.icr_min, "cash_flows": self.cash_flows,
        }


def _norm(s: object) -> str:
    return " ".join(str(s or "").split()).lower()


def _input_sheet(wb):
    for ws in wb.worksheets:
        if "исх проект" in _norm(ws.title):
            return ws
    raise ValueError("Не найден лист исходных проектных данных модели долгового риска")


def _row_values(ws, *keywords: str) -> list[float | None]:
    """Значения года 0–3 (столбцы D–G) строки, чей текст (B) содержит все ключевые слова."""
    keys = [k.lower() for k in keywords]
    for r in range(1, ws.max_row + 1):
        label = _norm(ws.cell(r, 2).value)
        if label and all(k in label for k in keys):
            return [ws.cell(r, c).value for c in range(4, 8)]
    return [None, None, None, None]


def _num(v) -> float | None:
    return float(v) if isinstance(v, (int, float)) else None


def _min_positive(vals: list[float | None]) -> float | None:
    xs = [float(v) for v in vals if isinstance(v, (int, float)) and v > 0]
    return min(xs) if xs else None


def parse_debt_risk_model(path: str | Path) -> ProjectFinance:
    """Извлечь финансовые показатели проекта из заполненного шаблона модели."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = _input_sheet(wb)

    capex = [float(v) for v in _row_values(ws, "капитальные затраты", "capex") if isinstance(v, (int, float))]
    dscr = _row_values(ws, "коэффициент покрытия задолженности")
    icr = _row_values(ws, "коэффициент покрытия процентов")
    cf = _row_values(ws, "денежный поток")
    dscr_stress = _row_values(ws, "стресс коэффициент покрытия задолженности")

    return ProjectFinance(
        capex_by_year=capex,
        capex_total=round(sum(capex), 4),
        npv=_num(_row_values(ws, "npv")[0]),
        irr=_num(_row_values(ws, "внутренняя норма доходности", "irr")[0]),
        npv_stress=_num(_row_values(ws, "стресс npv")[0]),
        irr_stress=_num(_row_values(ws, "норма доходности", "стресс")[0]),
        dscr_by_year=[_num(v) for v in dscr],
        dscr_min=_min_positive(dscr),
        dscr_stress_min=_min_positive(dscr_stress),
        icr_by_year=[_num(v) for v in icr],
        icr_min=_min_positive(icr),
        cash_flows=[_num(v) for v in cf],
    )


def credit_limit_from_finance(pf: ProjectFinance, dscr_norm: float = DSCR_NORM_DEFAULT) -> float | None:
    """Предел финансирования по кредитоспособности (граница долгового риска).

    Если минимальный DSCR проекта не ниже норматива — ограничение не накладывается
    (None). Иначе финансируемая сумма ограничивается долей запаса обслуживания долга:
    credit_limit = CAPEX · (DSCR_min / норматив).
    """
    if pf.dscr_min is None or pf.dscr_min >= dscr_norm:
        return None
    return round(pf.capex_total * max(pf.dscr_min, 0.0) / dscr_norm, 2)


def project_finance_to_opt(
    pf: ProjectFinance, name: str, *, score: float = 0.0,
    var_type: str = "continuous", min_share: float = 0.5,
    dscr_norm: float = DSCR_NORM_DEFAULT, use_stress: bool = False, key: str | None = None,
) -> OptProject:
    """Преобразовать финансовые показатели в кандидата оптимизации."""
    effect = pf.npv_stress if use_stress else pf.npv
    return OptProject(
        key=key or name, name=name,
        cost=pf.capex_total,
        effect=float(effect or 0.0),
        var_type=var_type, score=score, min_share=min_share,
        credit_limit=credit_limit_from_finance(pf, dscr_norm),
    )
